"""Milestone M2: visual spot-check of the material classification.

A person compares each sampled point with high-resolution imagery (Google Maps satellite view)
and records what is really there; agreement is the validation figure in the report.

Sampling: stratified, `per_class` points for each material class, inside the display
district, only where the whole 3 x 3 neighbourhood (30 m) has the same class, so a small
geolocation offset between Sentinel-2 and the imagery cannot move a point onto another class.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import xarray as xr
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pyproj import Transformer
from scipy.ndimage import maximum_filter, minimum_filter

from coolcairo.classify import Material
from coolcairo.config import Config, projected_bbox

# What a person looking at a satellite photo should call each class.
CLASS_LABELS = {
    Material.VEGETATION: "Vegetation",
    Material.DARK: "Dark surface",
    Material.BRIGHT: "Bright surface",
    Material.SOIL: "Bare soil / sand",
}
CLASS_HELP = {
    "Vegetation": "Trees, grass, gardens, green fields",
    "Dark surface": "Asphalt road, dark roof, dark paving",
    "Bright surface": "Pale concrete roof, light paving, white/beige building tops",
    "Bare soil / sand": "Empty sandy or dusty land, unbuilt plots, desert",
}
ANSWERS = [*CLASS_HELP, "Unsure / can't tell"]


def sample_points(
    cfg: Config, material: xr.DataArray, per_class: int = 5, seed: int = 42
) -> pd.DataFrame:
    """Stratified random points (pixel centres) in homogeneous 30 m patches of each class."""
    min_x, min_y, max_x, max_y = projected_bbox(
        cfg["display_aoi"]["bbox_wgs84"], cfg.crs, cfg.block_size
    )
    m = material.sel(x=slice(min_x, max_x), y=slice(max_y, min_y))
    values = m.values
    uniform = maximum_filter(values, size=3) == minimum_filter(values, size=3)

    rng = np.random.default_rng(seed)
    rows = []
    for cls in CLASS_LABELS:
        yy, xx = np.nonzero(uniform & (values == cls))
        take = rng.choice(len(yy), size=min(per_class, len(yy)), replace=False)
        for k in take:
            rows.append({"x": float(m.x[xx[k]]), "y": float(m.y[yy[k]]), "predicted": cls})
    points = pd.DataFrame(rows).sample(frac=1, random_state=seed).reset_index(drop=True)

    to_wgs = Transformer.from_crs(cfg.crs, "EPSG:4326", always_xy=True)
    points["lon"], points["lat"] = to_wgs.transform(points.x.values, points.y.values)
    points["predicted_label"] = points.predicted.map(lambda c: CLASS_LABELS[Material(c)])
    return points


def maps_link(lat: float, lon: float) -> str:
    """Google Maps satellite view centred on the point, zoomed to roughly building level."""
    return f"https://www.google.com/maps/@{lat:.6f},{lon:.6f},80m/data=!3m1!1e3"


FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # Cells the checker fills in.
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_sheet(points: pd.DataFrame, path: Path) -> Path:
    wb = Workbook()
    _instructions(wb.active)
    ws = wb.create_sheet("Spot check")
    headers = ["#", "Latitude", "Longitude", "Open in Google Maps", "Pipeline says",
               "What you see (choose)", "Match?", "Notes (optional)"]
    widths = [5, 11, 11, 22, 18, 26, 9, 40]
    for col, (h, w) in enumerate(zip(headers, widths, strict=True), start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30

    answers = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"',
                             allow_blank=True, showDropDown=False)
    ws.add_data_validation(answers)

    for i, p in points.iterrows():
        r = i + 2
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=round(p.lat, 6))
        ws.cell(row=r, column=3, value=round(p.lon, 6))
        link = ws.cell(row=r, column=4, value="Open satellite view")
        link.hyperlink = maps_link(p.lat, p.lon)
        ws.cell(row=r, column=5, value=p.predicted_label)
        answer = ws.cell(row=r, column=6)
        answer.fill = INPUT_FILL
        answers.add(answer)
        ws.cell(row=r, column=7,
                value=f'=IF(OR(F{r}="",F{r}="Unsure / can\'t tell"),"",IF(F{r}=E{r},"Yes","No"))')
        ws.cell(row=r, column=8).fill = INPUT_FILL
        for col in range(1, 9):
            cell = ws.cell(row=r, column=col)
            cell.border = BOX
            cell.font = Font(name=FONT, color="0563C1", underline="single") if col == 4 \
                else Font(name=FONT)
            if col in (1, 7):
                cell.alignment = Alignment(horizontal="center")

    last = len(points) + 1
    s = last + 2
    # Summary: label merged across A:E (so it is never cut off), value in F.
    summary = [
        ("Points checked (excluding Unsure)",
         f'=COUNTIF(G2:G{last},"Yes")+COUNTIF(G2:G{last},"No")'),
        ("Matches", f'=COUNTIF(G2:G{last},"Yes")'),
        ("Agreement", f'=IF(F{s}=0,"",F{s + 1}/F{s})'),
    ]
    for k, (label, formula) in enumerate(summary):
        r = s + k
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = Font(name=FONT, bold=True)
        lc.alignment = Alignment(horizontal="right")
        vc = ws.cell(row=r, column=6, value=formula)
        vc.font = Font(name=FONT, bold=True)
        vc.alignment = Alignment(horizontal="center")
    ws.cell(row=s + 2, column=6).number_format = "0%"
    ws.cell(row=s, column=6).comment = Comment(
        "Rows marked Unsure are left out of the agreement figure.", "CoolCairo")
    ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _instructions(ws: object) -> None:
    ws.title = "How to fill in"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 70
    lines = [
        ("CoolCairo: M2 spot check", None),
        ("", None),
        ("What this is", "Our pipeline labels every 10 m patch of Nasr City from satellite data. "
         "This sheet checks 20 random points by eye against high-resolution imagery."),
        ("Time needed", "About 20-30 minutes. No technical knowledge needed."),
        ("Steps", "1. Go to the 'Spot check' tab.\n"
         "2. For each row, click 'Open satellite view' (opens Google Maps at the point).\n"
         "3. Look at what is exactly at the centre of the map.\n"
         "4. Pick what you see from the dropdown in the yellow column. Do not look at the "
         "'Pipeline says' column first, to avoid bias.\n"
         "5. If you cannot tell, choose 'Unsure / can't tell' and add a note."),
        ("Only edit", "The yellow cells: 'What you see' and 'Notes'. Everything else is filled "
         "in automatically."),
        ("", None),
        ("Categories", None),
        *[(k, v) for k, v in CLASS_HELP.items()],
        ("Unsure / can't tell", "Clouds, blurry imagery, or a point exactly on an edge"),
        ("", None),
        ("Example (not counted)", "Point on a wide asphalt road -> choose 'Dark surface'; "
         "note: 'road, Abbas El Akkad St'."),
    ]
    for r, (a, b) in enumerate(lines, start=1):
        ca = ws.cell(row=r, column=1, value=a)
        ca.font = Font(name=FONT, bold=True, size=14 if r == 1 else 11)
        ca.alignment = Alignment(vertical="top")
        if b is not None:
            cb = ws.cell(row=r, column=2, value=b)
            cb.font = Font(name=FONT)
            cb.alignment = Alignment(wrap_text=True, vertical="top")


def read_results(path: Path) -> pd.DataFrame:
    """Filled sheet -> one row per point with predicted, observed and match."""
    ws = load_workbook(path, data_only=False)["Spot check"]
    rows = []
    for r in range(2, ws.max_row + 1):
        num, predicted, seen = ws.cell(r, 1).value, ws.cell(r, 5).value, ws.cell(r, 6).value
        if not isinstance(num, int):
            break
        rows.append({"point": num, "predicted": predicted, "observed": seen})
    df = pd.DataFrame(rows)
    usable = df.observed.notna() & (df.observed != ANSWERS[-1])
    df["match"] = np.where(usable, df.predicted == df.observed, np.nan)
    return df
